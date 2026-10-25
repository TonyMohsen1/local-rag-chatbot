import mimetypes

mimetypes.add_type("application/javascript", ".js")
mimetypes.add_type("text/css", ".css")

import hashlib

import chromadb
import ollama
import streamlit as st
from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

LLM_MODEL = "qwen2.5:3b"
EMBED_MODEL = "bge-m3"
CHUNK_SIZE = 800
CHUNK_OVERLAP = 100
TOP_K = 4
DB_PATH = "./chroma_db"
COLLECTION = "docs"

splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
    separators=["\n\n", "\n", ".", "،", "؟", "!", " ", ""],
)


@st.cache_resource
def get_collection():
    client = chromadb.PersistentClient(path=DB_PATH)
    return client.get_or_create_collection(
        COLLECTION, metadata={"hnsw:space": "cosine"}
    )


def extract_pages(file) -> list[tuple[int, str]]:
    name = file.name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(file)
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    if name.endswith(".docx"):
        doc = Document(file)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return [(1, "\n".join(parts))]
    return []


def embed(texts: list[str]) -> list[list[float]]:
    out = []
    for i in range(0, len(texts), 32):
        out.extend(ollama.embed(model=EMBED_MODEL, input=texts[i : i + 32])["embeddings"])
    return out


def index_file(file) -> str:
    data = file.getvalue()
    file_hash = hashlib.md5(data).hexdigest()[:12]
    col = get_collection()

    if col.get(where={"file_hash": file_hash}, limit=1)["ids"]:
        return "exists"

    pages = extract_pages(file)
    texts, metas, ids = [], [], []
    for page_no, text in pages:
        for j, chunk in enumerate(splitter.split_text(text)):
            texts.append(chunk)
            metas.append({"source": file.name, "page": page_no, "file_hash": file_hash})
            ids.append(f"{file_hash}_{page_no}_{j}")

    if not texts:
        return "empty"

    col.add(ids=ids, documents=texts, metadatas=metas, embeddings=embed(texts))
    return "ok"


def retrieve(question: str):
    col = get_collection()
    if col.count() == 0:
        return []
    q_emb = embed([question])[0]
    res = col.query(query_embeddings=[q_emb], n_results=min(TOP_K, col.count()))
    return list(zip(res["documents"][0], res["metadatas"][0]))


def build_prompt(question: str, hits) -> str:
    context = "\n\n".join(
        f"[المصدر: {m['source']} - ص{m['page']}]\n{d}" for d, m in hits
    )
    return (
        "أنت مساعد بيجاوب على الأسئلة اعتماداً على النصوص المرفقة فقط.\n"
        "- جاوب بنفس لغة السؤال.\n"
        "- لو الإجابة مش موجودة في النصوص، قول إنك مش لاقي الإجابة في الملفات.\n"
        "- اذكر المصدر ورقم الصفحة بين أقواس.\n\n"
        f"النصوص:\n{context}\n\nالسؤال: {question}"
    )


def stream_answer(prompt: str):
    for chunk in ollama.chat(
        model=LLM_MODEL,
        messages=[{"role": "user", "content": prompt}],
        stream=True,
    ):
        yield chunk["message"]["content"]


st.set_page_config(page_title="RAG Chatbot", page_icon="📄")
st.title("📄 اسأل ملفاتك")

with st.sidebar:
    st.header("الملفات")
    uploads = st.file_uploader(
        "ارفع PDF أو DOCX", type=["pdf", "docx"], accept_multiple_files=True
    )
    if uploads and st.button("📥 فهرسة الملفات", use_container_width=True):
        for f in uploads:
            with st.spinner(f"بفهرس {f.name}..."):
                status = index_file(f)
            if status == "ok":
                st.success(f"تمت فهرسة {f.name}")
            elif status == "exists":
                st.info(f"{f.name} متفهرس قبل كده")
            else:
                st.warning(
                    f"{f.name}: مفيش نص اتقرأ. لو PDF ممسوح ضوئياً "
                    "استخدم `ocrmypdf` الأول."
                )

    col = get_collection()
    files = {m["source"] for m in col.get()["metadatas"]} if col.count() else set()
    st.caption(f"الملفات المفهرسة: {len(files)} | الأجزاء: {col.count()}")
    for name in sorted(files):
        st.write(f"• {name}")

    if col.count() and st.button("🗑️ مسح كل البيانات", use_container_width=True):
        col.delete(ids=col.get()["ids"])
        st.session_state.messages = []
        st.rerun()

if "messages" not in st.session_state:
    st.session_state.messages = []

for msg in st.session_state.messages:
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])
        if msg.get("sources"):
            with st.expander("المصادر"):
                for s in msg["sources"]:
                    st.caption(s)

if question := st.chat_input("اكتب سؤالك عن الملفات..."):
    st.session_state.messages.append({"role": "user", "content": question})
    with st.chat_message("user"):
        st.markdown(question)

    with st.chat_message("assistant"):
        hits = retrieve(question)
        if not hits:
            answer, sources = "ارفع وافهرس ملف الأول عشان أقدر أجاوب.", []
            st.markdown(answer)
        else:
            answer = st.write_stream(stream_answer(build_prompt(question, hits)))
            sources = [
                f"{m['source']} — ص{m['page']}: {d[:150]}..." for d, m in hits
            ]
            with st.expander("المصادر"):
                for s in sources:
                    st.caption(s)

    st.session_state.messages.append(
        {"role": "assistant", "content": answer, "sources": sources}
    )
